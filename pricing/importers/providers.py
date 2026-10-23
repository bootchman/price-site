"""Разборщики прайсов текущих операторов (шаблоны). Новые форматы — через мастер загрузки (следующий шаг)."""
import json
import re
import urllib.request
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from .base import Importer, money
from pricing.reference import DESTINATION_ALIASES, US_STATES, terminal_code


def display_name(raw: str, state: str = "") -> str:
    n = str(raw or "")
    n = re.sub(r"\b\d{5}(-\d{4})?\b", " ", n)
    n = re.sub(r"\b[A-Za-z]\d[A-Za-z][ -]?\d[A-Za-z]\d\b", " ", n)
    n = re.sub(r"\(\s*\)", " ", n)
    if state:
        n = re.sub(r"\(\s*" + re.escape(state) + r"\s*\)", " ", n, flags=re.I)
        n = re.sub(r"\s" + re.escape(state) + r"\s*\)", " ", n, flags=re.I)
        n = re.sub(r",\s*" + re.escape(state) + r"\s*$", " ", n.strip(), flags=re.I)
        n = re.sub(r"\s" + re.escape(state) + r"\s*$", " ", n.strip())
    return re.sub(r"\s+", " ", n).strip(" ,")


# ----------------------------------------------------------------------------- Cargoloop
CARGOLOOP_URL = "https://www.cargoloop.com/LocalTransportation.aspx?e6d7dfd5="
CARGOLOOP_TERMINALS = ["Savannah, GA", "Elizabeth, NJ", "Houston, TX", "Los Angeles, CA", "Indianapolis, IN", "Montreal, QC"]
# Коэффициенты Cargoloop к базовой цене (база — седан, SUV стандартный, фургон стандартный)
CARGOLOOP_MULTIPLIERS = {"sedan": 1.0, "suv": 1.0, "suv_big": 1.5, "pickup": 1.5, "moto": 0.8}


class CargoloopImporter(Importer):
    provider_name = "Cargoloop"
    allow_city_match = False

    def parse(self, html: str):
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(html, "html.parser")
        text = soup.get_text(" ", strip=True)
        m = re.search(r"Valid from ([A-Za-z]+ \d{1,2})(?:st|nd|rd|th)?,? (\d{4})", text)
        if m:
            try:
                self.meta["valid_from"] = datetime.strptime(f"{m.group(1)} {m.group(2)}", "%B %d %Y").date()
            except ValueError:
                pass
        table = max(soup.find_all("table"), key=lambda t: len(t.find_all("tr")))
        header = [c.get_text(strip=True) for c in table.find("tr").find_all(["td", "th"])]
        if header[5:11] != CARGOLOOP_TERMINALS:
            self.warnings.append(f"Колонки складов изменились: {header[5:]}")
        for tr in table.find_all("tr")[1:]:
            td = tr.find_all("td")
            if len(td) < 11:
                continue
            name, auction, city, state, zip_code = (c.get_text(strip=True) for c in td[:5])
            yard = self.yard(auction, state, display_name(name, state), city=city.title(), zip_code=zip_code)
            self.counters["строк прайса"] += 1
            for i, term_name in enumerate(CARGOLOOP_TERMINALS):
                cell = td[5 + i]
                price = money(cell.get_text(strip=True))
                if price is None:
                    continue
                preferred = "background-color" in (cell.get("style") or "")
                for vtype, k in CARGOLOOP_MULTIPLIERS.items():
                    self.land(yard, terminal_code(term_name), vtype, price * k, preferred=preferred,
                              estimate=(k != 1.0), note="" if k == 1.0 else f"база × {k}")

    def load(self, source: str):
        if source.startswith("http"):
            req = urllib.request.Request(source, headers={"User-Agent": "Mozilla/5.0"})
            html = urllib.request.urlopen(req, timeout=60).read().decode("utf-8", "ignore")
        else:
            html = Path(source).read_text(encoding="utf-8", errors="ignore")
        self.parse(html)


# ----------------------------------------------------------------------------- Atlantic
ATLANTIC_TYPES = {"car": "sedan", "suv": "suv", "pickup": "pickup", "motorcycle": "moto"}


class AtlanticImporter(Importer):
    provider_name = "Atlantic"

    def load(self, source: str):
        import openpyxl

        ws = openpyxl.load_workbook(source, read_only=True, data_only=True).active
        rows = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(rows)]
        if header[:4] != ["from", "to", "cargo type", "price"]:
            self.warnings.append(f"Неожиданные колонки: {header}")
        cars = {}
        explicit = defaultdict(set)
        for frm, to, cargo, price in (r[:4] for r in rows):
            self.counters["строк прайса"] += 1
            parts = [p.strip() for p in str(frm or "").split(" - ")]
            if len(parts) < 3 or not parts[0]:
                self.counters["битых строк «From»"] += 1
                continue
            auction, state_full = parts[-1], parts[-2]
            state = US_STATES.get(state_full)
            if not state:
                self.counters["неизвестный штат"] += 1
                continue
            extra = [p for p in parts[1:-2] if not re.fullmatch(r"[A-Z]{2}", p) and p.lower() != auction.lower()]
            name = " ".join([parts[0]] + extra)
            term = terminal_code(to)
            vtype = ATLANTIC_TYPES.get(str(cargo or "").strip().lower())
            if not term or not vtype:
                self.counters["пропущено (склад или тип груза)"] += 1
                continue
            value = money(price)
            if not value:
                self.counters["нет сервиса (цена 0)"] += 1
                continue
            yard = self.yard(auction, state, name)
            self.land(yard, term, vtype, value)
            explicit[(yard.id if yard else None, term)].add(vtype)
            if vtype == "sedan":
                cars[(yard.id if yard else None, term)] = (yard, value)
        # SUV без отдельной цены считаем как седан (у Atlantic SUV/седан ≈ 1.00), с пометкой «оценка»
        for key, (yard, value) in cars.items():
            if "suv" not in explicit[key]:
                self.land(yard, key[1], "suv", value, estimate=True, note="как седан")


# ----------------------------------------------------------------------------- VS Brothers (калькулятор)
VS_SHEETS = {"Sedan": "sedan", "Small, Medium - SUV": "suv", "Big - SUV": "suv_big", "Pickup": "pickup", "Motorcyrcles": "moto"}
VS_COLUMNS = {"ca california": "LAX", "ga georgia": "SAV", "nj new jersey": "ELZ", "tx texas": "HOU"}
VS_SEA_SEDAN_DEFAULT = {"LAX": 900, "SAV": 500, "ELZ": 550, "HOU": 650}  # море VS для седана (Alex, 23.10.2026)


class VSCalculatorImporter(Importer):
    """«Суша + море» до порта: суша = седан − море седана (одинакова для всех типов),
    море по типу = цена типа − суша (надбавки по типу — морская часть)."""

    provider_name = "VS Brothers"

    def __init__(self, *args, sea_sedan=None, **kwargs):
        kwargs.setdefault("is_bundle", True)
        super().__init__(*args, **kwargs)
        self.sea_sedan = sea_sedan or VS_SEA_SEDAN_DEFAULT

    def load(self, source: str):
        import openpyxl

        wb = openpyxl.load_workbook(source, data_only=True)
        title = str(wb.worksheets[0]["A1"].value or "")
        m = re.match(r"(.+?) to (.+?),\s*(.+)$", title)
        if not m:
            raise ValueError(f"Не распознан заголовок листа: {title!r}")
        auction = "IAAI" if m.group(1).replace(" ", "").upper() == "IAAI" else m.group(1).strip().title()
        port = DESTINATION_ALIASES.get(m.group(3).strip().lower())
        if not port:
            raise ValueError(f"Неизвестный порт назначения: {m.group(3)!r}")
        self.meta["scope"] = auction.lower()

        sheets = {}
        for ws in wb.worksheets:
            vtype = VS_SHEETS.get(ws.title.strip())
            if not vtype:
                self.warnings.append(f"Пропущен лист {ws.title!r}")
                continue
            header = [str(c or "").strip().lower() for c in next(ws.iter_rows(min_row=3, max_row=3, values_only=True))]
            cols = {i: VS_COLUMNS[h] for i, h in enumerate(header) if h in VS_COLUMNS}
            data = {}
            for row in ws.iter_rows(min_row=4, values_only=True):
                loc = str(row[0] or "").strip()
                mm = re.match(r"^([A-Z]{2})\s*-\s*(.+)$", loc)
                if not mm:
                    continue
                data[(mm.group(1), mm.group(2).strip())] = {t: money(row[i]) for i, t in cols.items()}
            sheets[vtype] = data

        sedan = sheets["sedan"]
        sea_values = defaultdict(Counter)
        for (state, name), prices in sedan.items():
            self.counters["площадок в прайсе"] += 1
            yard = self.yard(auction, state, name)
            for term, full in prices.items():
                if full is None:
                    continue
                land = full - self.sea_sedan[term]
                for vtype, data in sheets.items():
                    type_full = (data.get((state, name)) or {}).get(term)
                    if type_full is None:
                        continue
                    self.land(yard, term, vtype, land, note="полная цена − море")
                    sea_values[(term, vtype)][round(type_full - land, 2)] += 1
        for (term, vtype), values in sea_values.items():
            value, count = values.most_common(1)[0]
            if len(values) > 1:
                self.warnings.append(f"Море {term}/{vtype}: разные значения {dict(values)}, взято {value}")
            self.sea(term, port, vtype, value, note="из калькулятора VS")


# ----------------------------------------------------------------------------- Shipper 101 (Google Таблица)
class Shipper101Importer(Importer):
    """JSON со строками листа «Location Copart/IAAI» (выгружает root-скрипт из Google Таблицы)."""

    provider_name = "Shipper 101"

    def load(self, source: str, auction: str):
        rows = json.loads(Path(source).read_text(encoding="utf-8"))
        self.meta["scope"] = auction.lower()
        for row in rows[1:]:
            row = list(row) + [""] * 4
            m = re.match(r"USA\s*-\s*([A-Z]{2})", str(row[0]).strip())
            price = money(row[3])
            if not m or price is None:
                continue
            self.counters["строк прайса"] += 1
            term = terminal_code(row[2])
            if not term:
                self.counters["неизвестный склад"] += 1
                continue
            yard = self.yard(auction, m.group(1), str(row[1]).strip())
            self.land(yard, term, "sedan", price)
            # Цена одна на машину; для SUV малого/среднего берём ту же с пометкой «оценка»
            self.land(yard, term, "suv", price, estimate=True, note="как седан")


# ----------------------------------------------------------------------------- AML (Google Таблица «Порівняння»)
AML_PORTS = {"NJ": "ELZ", "GA": "SAV", "FL": "MIA", "TX": "HOU", "CA": "LAX", "WA": "SEA"}


class AMLImporter(Importer):
    """Лист Copart/IAAI: блок «Давид (прайс суша-море)» = AML (Порт, Суша, Море, Послуги, Разом).
    Лист «Ставки»: море AML по портам и «Послуги — AML». Левый блок (Anna) — не оператор, не загружаем."""

    provider_name = "AML"

    def load(self, sheet_json: str, rates_json: str, auction: str):
        self.meta["scope"] = auction.lower()
        rows = json.loads(Path(sheet_json).read_text(encoding="utf-8"))
        header_idx = next(i for i, r in enumerate(rows) if r[:2] == ["Штат", "Локація"])
        header = rows[header_idx]
        aml_from = next(i for i, h in enumerate(header) if i > 2 and h == "Порт" and i >= 8)
        if header[aml_from:aml_from + 4] != ["Порт", "Суша", "Море", "Послуги"]:
            self.warnings.append(f"Неожиданные колонки AML: {header[aml_from:aml_from + 5]}")
        for row in rows[header_idx + 1:]:
            row = list(row) + [""] * 16
            state, loc = str(row[0]).strip(), str(row[1]).strip()
            port, land = str(row[aml_from]).strip().upper(), money(row[aml_from + 1])
            if not state or not loc:
                continue
            self.counters["строк прайса"] += 1
            if land is None:
                self.counters["нет цены суши"] += 1
                continue
            term = AML_PORTS.get(port)
            if not term:
                self.counters["неизвестный порт"] += 1
                continue
            yard = self.yard(auction, state, loc)
            self.land(yard, term, "sedan", land)
            self.land(yard, term, "suv", land, estimate=True, note="как седан")

        rates = json.loads(Path(rates_json).read_text(encoding="utf-8"))
        hdr = next((r for r in rates if r and r[0] == "Порт США"), [])
        col = next((i for i, h in enumerate(hdr) if "Давид" in h or "AML" in h), None)
        for r in rates:
            r = list(r) + [""] * 4
            if col and r[0] in AML_PORTS:
                sea = money(r[col])
                if sea is not None:
                    self.sea(AML_PORTS[r[0]], "KLJ", "sedan", sea, note="лист «Ставки»")
                    self.sea(AML_PORTS[r[0]], "KLJ", "suv", sea, estimate=True, note="как седан")
            if str(r[0]).strip().lower() in ("послуги — aml", "послуги - aml"):
                self.charge("services", "Послуги", money(r[1]), note="лист «Ставки»")
