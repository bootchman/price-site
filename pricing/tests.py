import json
import tempfile
from decimal import Decimal
from pathlib import Path

import openpyxl
from django.test import TestCase

from pricing.importers import providers
from pricing.models import DestinationPort, PriceList, Rate, VehicleType, Yard
from pricing.reference import name_key
from pricing.routes import find_routes

TMP = Path(tempfile.mkdtemp())

CARGOLOOP_HTML = """
<html><body><p>Valid from October 1st, 2026 Local Transportation Prices</p>
<table><tr><td>x</td></tr></table>
<table>
<tr><th>Auction Location</th><th>Auction</th><th>City</th><th>State</th><th>Zip</th>
<th>Savannah, GA</th><th>Elizabeth, NJ</th><th>Houston, TX</th><th>Los Angeles, CA</th><th>Indianapolis, IN</th><th>Montreal, QC</th></tr>
<tr><td>MOBILE SOUTH (AL) 36582</td><td>Copart</td><td>THEODORE</td><td>AL</td><td>36582</td>
<td style="background-color: #47d78d">$500</td><td>$900</td><td>$700</td><td>$1,500</td><td>$800</td><td></td></tr>
</table></body></html>
"""


def write_vs(path, title, sheets):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        ws["A1"] = title
        ws.append([])
        ws.append(["Auction Location", "CA California", "GA Georgia", "NJ New Jersey", "TX Texas"])
        for row in rows:
            ws.append(row)
    wb.save(path)


class NameKeyTests(TestCase):
    def test_variants_collapse(self):
        self.assertEqual(name_key("MOBILE SOUTH (AL) 36582", "AL"), "MOBILE SOUTH")
        self.assertEqual(name_key("Mobile South", "AL"), "MOBILE SOUTH")
        self.assertEqual(name_key("ACE - Carson (CA) 90248", "CA"), name_key("ACE Carson", "CA"))
        self.assertEqual(name_key("ABILENE, TX 79601", "TX"), "ABILENE")
        self.assertEqual(name_key("ABBOTSFORD, BC V4X-1G8", "BC"), "ABBOTSFORD")
        self.assertEqual(name_key("TUCSON AZ) 85706", "AZ"), "TUCSON")
        self.assertEqual(name_key("CONCORD NC ( )", "NC"), "CONCORD")
        self.assertEqual(providers.display_name("TUCSON AZ) 85706", "AZ"), "TUCSON")
        self.assertEqual(providers.display_name("CONCORD NC ( )", "NC"), "CONCORD")


class ImportAndRouteTests(TestCase):
    def import_cargoloop(self):
        imp = providers.CargoloopImporter("t", "test")
        imp.parse(CARGOLOOP_HTML)
        return imp.save()

    def import_vs(self):
        path = TMP / "vs.xlsx"
        write_vs(path, "COPART to Lithuania,Klaipeda", {
            "Sedan": [["AL - MOBILE SOUTH", "-", 950, 1400, "-"]],
            "Pickup": [["AL - MOBILE SOUTH", "-", 1350, 1800, "-"]],
        })
        imp = providers.VSCalculatorImporter("t", str(path))
        imp.load(str(path))
        return imp.save()

    def test_cargoloop(self):
        pl = self.import_cargoloop()
        self.assertEqual(str(pl.valid_from), "2026-10-01")
        yard = Yard.objects.get()
        self.assertEqual((yard.name, yard.zip, yard.city), ("MOBILE SOUTH", "36582", "Theodore"))
        sav = Rate.objects.get(stage="land", to_terminal__code="SAV", vehicle_type__code="sedan")
        self.assertEqual((sav.amount, sav.is_preferred), (Decimal("500"), True))
        moto = Rate.objects.get(stage="land", to_terminal__code="SAV", vehicle_type__code="moto")
        self.assertEqual((moto.amount, moto.is_estimate), (Decimal("400"), True))
        self.assertFalse(Rate.objects.filter(to_terminal__code="MTL").exists())

    def test_vs_splits_bundle_and_matches_yard(self):
        self.import_cargoloop()
        self.import_vs()
        self.assertEqual(Yard.objects.count(), 1)  # «AL - MOBILE SOUTH» VS = площадка Cargoloop
        land = Rate.objects.get(stage="land", price_list__provider__name="VS Brothers", to_terminal__code="SAV",
                                vehicle_type__code="pickup")
        self.assertEqual(land.amount, Decimal("450"))  # 950 − 500, суша одинакова для всех типов
        sea = Rate.objects.get(stage="sea", from_terminal__code="SAV", vehicle_type__code="pickup")
        self.assertEqual(sea.amount, Decimal("900"))  # 1350 − 450 → надбавка 400 в море

    def test_route_combines_providers(self):
        self.import_cargoloop()
        self.import_vs()
        yard = Yard.objects.get()
        sedan = VehicleType.objects.get(code="sedan")
        klj = DestinationPort.objects.get(code="KLJ")
        full, land_only = find_routes(yard, sedan, klj)
        totals = [(o.land.price_list.provider.name, o.land.to_terminal.code, int(o.total), o.same_provider) for o in full]
        # Savannah: суша VS 450 / Cargoloop 500 + море VS 500; Elizabeth: VS 850 / Cargoloop 900 + море 550
        self.assertEqual(totals[0], ("VS Brothers", "SAV", 950, True))
        self.assertIn(("Cargoloop", "SAV", 1000, False), totals)
        self.assertIn(("Cargoloop", "ELZ", 1450, False), totals)
        self.assertEqual({r.to_terminal.code for r in land_only}, {"HOU", "LAX", "IND"})

    def test_sea_from_two_vs_files_not_duplicated(self):
        self.import_cargoloop()
        self.import_vs()
        path = TMP / "vs_iaai.xlsx"
        write_vs(path, "I A A I to Lithuania,Klaipeda", {"Sedan": [["AL - Mobile South", "-", 950, 1400, "-"]]})
        imp = providers.VSCalculatorImporter("t", str(path))
        imp.load(str(path))
        imp.save()
        yard = Yard.objects.get(auction__code="Copart")
        full, _ = find_routes(yard, VehicleType.objects.get(code="sedan"), DestinationPort.objects.get(code="KLJ"))
        keys = [(o.land.id, o.sea.price_list.provider_id, o.sea.from_terminal_id) for o in full]
        self.assertEqual(len(keys), len(set(keys)))

    def test_new_version_archives_previous(self):
        first = self.import_cargoloop()
        second = self.import_cargoloop()
        first.refresh_from_db()
        self.assertEqual((first.status, second.status), (PriceList.ARCHIVED, PriceList.PUBLISHED))

    def test_atlantic_new_york_matches_elizabeth_sea(self):
        self.import_vs()
        path = TMP / "atl.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["From", "To", "Cargo Type", "Price"])
        ws.append(["MOBILE SOUTH - Alabama - Copart", "New York", "car", 300])
        ws.append(["MOBILE SOUTH - Alabama - Copart", "New York", "motorcycle", 250])
        ws.append(["Sioux Falls - SD - Adesa - South Dakota - Adesa", "Chicago", "car", 0])
        wb.save(path)
        imp = providers.AtlanticImporter("t", str(path))
        imp.load(str(path))
        pl = imp.save()
        self.assertEqual(pl.stats.get("нет сервиса (цена 0)"), 1)
        suv = Rate.objects.get(price_list=pl, vehicle_type__code="suv")
        self.assertTrue(suv.is_estimate)
        yard = Yard.objects.get(name_key="MOBILE SOUTH")
        full, _ = find_routes(yard, VehicleType.objects.get(code="sedan"), DestinationPort.objects.get(code="KLJ"))
        atl = [o for o in full if o.land.price_list.provider.name == "Atlantic"][0]
        self.assertEqual((int(atl.total), atl.sea.from_terminal.code), (850, "ELZ"))  # 300 + 550
        self.assertIn("New York ≈ Elizabeth, NJ", "; ".join(atl.notes))

    def test_city_and_prefix_fallback(self):
        self.import_cargoloop()  # MOBILE SOUTH, город Theodore
        path = TMP / "s101b.json"
        path.write_text(json.dumps([
            ["State", "Location, from", "Warehouse - Loading Port", "101"],
            ["USA - AL", "THEODORE", "Savannah, GA", "480"],
            ["USA - AL", "MOBILE", "Savannah, GA", "490"],
        ]), encoding="utf-8")
        imp = providers.Shipper101Importer("t", str(path))
        imp.load(str(path), "Copart")
        pl = imp.save()
        self.assertEqual(pl.stats.get("площадок найдено по городу"), 1)
        self.assertEqual(pl.stats.get("новых площадок"), 1)  # «MOBILE» ≠ «MOBILE SOUTH»
        self.assertEqual(Yard.objects.count(), 2)

    def test_shipper101(self):
        path = TMP / "s101.json"
        path.write_text(json.dumps([
            ["State", "Location, from", "Warehouse - Loading Port", "101"],
            ["USA - AL", "Mobile South", "Savannah, GA", "500"],
            ["USA - AL", "NEWTON", "Savannah, GA"],
        ]), encoding="utf-8")
        imp = providers.Shipper101Importer("t", str(path))
        imp.load(str(path), "Copart")
        imp.save()
        self.assertEqual(Rate.objects.filter(stage="land", vehicle_type__code="sedan").count(), 1)

    def test_route_page(self):
        from django.contrib.auth.models import User
        from accounts.models import TelegramProfile
        from unittest import mock

        self.import_cargoloop()
        self.import_vs()
        user = User.objects.create(username="tg1")
        TelegramProfile.objects.create(user=user, telegram_id=1, phone="380000000000")
        self.client.force_login(user)
        yard = Yard.objects.get()
        with mock.patch("accounts.middleware.is_allowed", return_value=True):
            resp = self.client.get("/", {"auction": yard.auction_id, "yard": yard.id, "vt": "sedan", "port": "KLJ"})
            self.assertContains(resp, "$950")
            self.assertContains(resp, "один оператор")
            self.assertEqual(self.client.get("/operators/").status_code, 200)
